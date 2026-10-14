#!/usr/bin/env python3
"""
Extract text from Word documents
"""

import sys
from docx import Document

def extract_word_text(doc_path):
    """Extract text from a Word document"""
    try:
        doc = Document(doc_path)
        full_text = []

        for paragraph in doc.paragraphs:
            full_text.append(paragraph.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text.append(cell.text)

        return '\n'.join(full_text)

    except Exception as e:
        return f"Error extracting Word document text: {str(e)}"

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Error: Document file path not provided")
        sys.exit(1)

    doc_path = sys.argv[1]
    extracted_text = extract_word_text(doc_path)
    print(extracted_text)